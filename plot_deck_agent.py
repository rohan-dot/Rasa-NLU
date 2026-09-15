#!/usr/bin/env python3
"""
plot_deck_agent.py — point it at a folder of plot images; a vision model (through the
LiteLLM gateway) reads each plot and writes reviewer-facing takeaways, then the script
assembles a 16:9 PowerPoint deck: title slide, key-findings slide, and one slide per
plot (plot left, takeaways right, plain-English description in the speaker notes).

    python plot_deck_agent.py eval_plots/ -o eval_deck.pptx --model <gateway opus alias> \\
        --title "FCG route planner — evaluation" \\
        --context "Geometry P/R/F1 vs Natural Earth ground truth; routing ablation; \\
                   miles-vs-clearances Pareto; judge agreement (Cohen's kappa)"

Environment (same names the rest of the FCG pipeline uses):
    LITELLM_BASE_URL   gateway URL (with or without /v1)
    LITELLM_API_KEY    gateway key
    AGENT_MODEL        default for --model
    AGENT_TLS_VERIFY   set to false for the self-signed cert

Per-plot analyses are cached in <out>.analysis.json keyed on image hash + model, so
re-runs only pay for new or changed plots. Edit the takeaways in that file and re-run
with --no-llm to rebuild the deck without any model calls.

pip install openai python-pptx
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import mimetypes
import os
import sys
import time
from datetime import date
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROMPT_VERSION = "1"        # bump to invalidate cached analyses when the prompts change
MAX_TOKENS = 4000           # generous so gateway aliases with extended thinking still fit
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}

PLOT_SYSTEM = """You analyse one evaluation plot at a time for a slide deck that technical reviewers will read without a presenter.

Rules:
- Closed world: state only what is visible in the plot. Read numbers off axes, labels and legends only when they are clearly legible; otherwise describe qualitatively ("roughly", "about"). Never invent values, sample sizes or methods that are not shown.
- No implementation detail: no code, file names, function names or tooling. Slides carry claims and design rationale only.
- Every takeaway is a complete sentence a reviewer could verify against the plot.

Return ONLY a JSON object (no markdown fences) with exactly this shape:
{
  "title": "slide title, at most 10 words, states the finding rather than the chart type",
  "plot_type": "e.g. Pareto curve, grouped bar chart, line plot",
  "takeaways": ["3 to 4 bullets, each at most 25 words"],
  "caveat": "one sentence on a limitation visible in the plot or on what it does not show, or an empty string",
  "notes": "2 to 4 plain-English sentences describing the plot, for speaker notes"
}"""

SYNTH_SYSTEM = """You write the 'Key findings' slide for a deck of evaluation plots. You receive the per-plot analyses as JSON.

Rules: only claims supported by those analyses; no implementation detail; no invented numbers; each finding traceable to a specific plot.

Return ONLY a JSON object (no markdown fences):
{
  "subtitle": "at most 12 words describing what the deck shows",
  "key_findings": ["3 to 5 bullets, each at most 25 words"]
}"""


# ----------------------------------------------------------------- LLM side

def log(msg: str) -> None:
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", file=sys.stderr, flush=True)


def make_client():
    import httpx
    from openai import OpenAI

    base_url = os.environ.get("LITELLM_BASE_URL")
    api_key = os.environ.get("LITELLM_API_KEY")
    if not base_url or not api_key:
        sys.exit("LITELLM_BASE_URL and LITELLM_API_KEY must be set")
    verify = os.environ.get("AGENT_TLS_VERIFY", "true").strip().lower() not in {"0", "false", "no"}
    return OpenAI(base_url=base_url, api_key=api_key,
                  http_client=httpx.Client(verify=verify, timeout=180.0))


def chat(client, model: str, system: str, user_content) -> str:
    resp = client.chat.completions.create(
        model=model,
        max_tokens=MAX_TOKENS,
        messages=[{"role": "system", "content": system},
                  {"role": "user", "content": user_content}],
    )
    return resp.choices[0].message.content or ""


def parse_json(text: str) -> dict:
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < 0:
        raise ValueError(f"no JSON object in model output: {text[:200]!r}")
    return json.loads(text[start:end + 1])


def ask_json(client, model: str, system: str, user_content) -> dict:
    """One retry with an explicit JSON-only nudge; anything else propagates to the caller."""
    last_err: Exception = ValueError("unreachable")
    for attempt in (1, 2):
        try:
            return parse_json(chat(client, model, system, user_content))
        except ValueError as e:  # includes json.JSONDecodeError
            last_err = e
            log(f"  malformed JSON (attempt {attempt}): {e}")
            blocks = user_content if isinstance(user_content, list) else [{"type": "text", "text": user_content}]
            user_content = blocks + [{"type": "text", "text": "Reply with the JSON object only."}]
    raise last_err


def image_block(path: Path) -> dict:
    mime = mimetypes.guess_type(path.name)[0] or "image/png"
    b64 = base64.b64encode(path.read_bytes()).decode("ascii")
    return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}


def analyse_plot(client, model: str, path: Path, context: str) -> dict:
    text = (f"Context for this set of plots: {context or 'none given; infer it from the plot itself'}\n"
            f"Image file name: {path.name}\n"
            "Analyse the attached plot.")
    result = ask_json(client, model, PLOT_SYSTEM, [{"type": "text", "text": text}, image_block(path)])
    result["title"] = str(result.get("title") or path.stem)
    result["takeaways"] = [str(t) for t in result.get("takeaways") or []][:5]
    result["caveat"] = str(result.get("caveat") or "")
    result["notes"] = str(result.get("notes") or "")
    return result


def synthesise(client, model: str, analyses: list, context: str) -> dict:
    payload = [{"file": name, **{k: a.get(k) for k in ("title", "plot_type", "takeaways", "caveat")}}
               for name, a in analyses]
    text = f"Context: {context or 'none given'}\n\nPer-plot analyses:\n{json.dumps(payload, indent=1)}"
    result = ask_json(client, model, SYNTH_SYSTEM, text)
    result["subtitle"] = str(result.get("subtitle") or "")
    result["key_findings"] = [str(k) for k in result.get("key_findings") or []][:6]
    return result


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def load_sidecar(path: Path) -> dict:
    if path.exists():
        try:
            return json.loads(path.read_text())
        except json.JSONDecodeError:
            log(f"{path} is not valid JSON; starting a fresh cache")
    return {"plots": {}, "synthesis": {}}


def save_sidecar(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, indent=2))


# ---------------------------------------------------------------- deck side

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
INK, MUTED = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80)


def _set_bullet(paragraph, indent=Inches(0.28)) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent)))
    pPr.set("indent", str(-int(indent)))
    etree.SubElement(pPr, qn("a:buChar"), char="•")


def add_text(slide, left, top, width, height, paragraphs, size=16, bold=False, color=INK,
             bullets=False, space_after=8):
    """paragraphs: list of str, or (str, overrides) tuples with per-paragraph overrides."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(paragraphs):
        text, opts = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if opts.get("bullet", bullets):
            _set_bullet(p)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        run.font.size = Pt(opts.get("size", size))
        run.font.bold = opts.get("bold", bold)
        run.font.italic = opts.get("italic", False)
        run.font.color.rgb = opts.get("color", color)
    return box


def fit_picture(slide, path: Path, left, top, box_w, box_h):
    pic = slide.shapes.add_picture(str(path), left, top, height=box_h)
    if pic.width > box_w:
        scale = box_w / pic.width
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = int(left + (box_w - pic.width) / 2)
    pic.top = int(top + (box_h - pic.height) / 2)
    return pic


def build_deck(out: Path, title: str, subtitle: str, findings: list, analyses: list, model: str) -> int:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    # 1. title slide
    s = prs.slides.add_slide(blank)
    add_text(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.4), [title], size=40, bold=True)
    add_text(s, Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.9),
             [subtitle or f"{len(analyses)} evaluation plots"], size=20, color=MUTED)
    add_text(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5),
             [date.today().strftime("%B %d, %Y")], size=12, color=MUTED)
    s.notes_slide.notes_text_frame.text = f"Generated by plot_deck_agent.py; plot analyses by {model}."

    # 2. key findings
    if findings:
        s = prs.slides.add_slide(blank)
        add_text(s, Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0), ["Key findings"], size=28, bold=True)
        add_text(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.4), findings,
                 size=18, bullets=True, space_after=12)

    # 3. one slide per plot
    for path, a in analyses:
        s = prs.slides.add_slide(blank)
        add_text(s, Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0), [a["title"]], size=26, bold=True)
        fit_picture(s, path, Inches(0.6), Inches(1.45), Inches(7.4), Inches(5.25))
        add_text(s, Inches(0.6), Inches(6.85), Inches(7.4), Inches(0.35), [path.name], size=9, color=MUTED)
        paras = [(t, {"bullet": True}) for t in a["takeaways"]]
        if a.get("caveat"):
            paras.append((a["caveat"], {"size": 12, "italic": True, "color": MUTED, "space_after": 0}))
        if paras:
            add_text(s, Inches(8.3), Inches(1.45), Inches(4.5), Inches(5.6), paras, size=16, space_after=10)
        s.notes_slide.notes_text_frame.text = f"{a.get('notes', '')}\n\nSource: {path}".strip()

    prs.save(str(out))
    return len(prs.slides)


# --------------------------------------------------------------------- main

def collect_images(inputs: list) -> list:
    paths = []
    for raw in inputs:
        p = Path(raw)
        if p.is_dir():
            paths += sorted(q for q in p.iterdir() if q.suffix.lower() in IMAGE_EXTS)
        elif p.suffix.lower() in IMAGE_EXTS:
            paths.append(p)
        else:
            log(f"skipping {p}: not a png/jpg (export matplotlib figures as png)")
    return paths


def placeholder(path: Path, reason: str) -> dict:
    return {"title": path.stem, "takeaways": [reason] if reason else [], "caveat": "", "notes": "", "error": True}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", help="plot image files and/or directories")
    ap.add_argument("-o", "--out", default="plot_deck.pptx")
    ap.add_argument("--model", default=os.environ.get("AGENT_MODEL"),
                    help="gateway model alias (default: $AGENT_MODEL)")
    ap.add_argument("--title", default="Evaluation results")
    ap.add_argument("--subtitle", default=None, help="override the model-written subtitle")
    ap.add_argument("--context", default="", help="one or two sentences on what the plots measure")
    ap.add_argument("--limit", type=int, default=None, help="only the first N plots (cheap trial run)")
    ap.add_argument("--no-cache", action="store_true", help="re-analyse every plot")
    ap.add_argument("--no-llm", action="store_true",
                    help="never call the model; build from cached analyses only")
    args = ap.parse_args()

    images = collect_images(args.inputs)[: args.limit]
    if not images:
        sys.exit("no png/jpg plots found")
    out = Path(args.out)
    sidecar_path = out.with_suffix(".analysis.json")
    sidecar = load_sidecar(sidecar_path)

    client = None
    if not args.no_llm:
        if not args.model:
            sys.exit("pass --model or set AGENT_MODEL to the gateway's alias for the model you want")
        client = make_client()
    model = args.model or "cached"

    analyses, failures = [], 0
    for i, path in enumerate(images, 1):
        h = file_hash(path)
        cached = sidecar["plots"].get(path.name) or {}
        fresh = cached.get("hash") == h and (
            args.no_llm or (cached.get("model") == model and cached.get("prompt") == PROMPT_VERSION))
        if fresh and (args.no_llm or not args.no_cache):
            log(f"[{i}/{len(images)}] {path.name}: cached")
            analyses.append((path, cached["analysis"]))
            continue
        if client is None:
            log(f"[{i}/{len(images)}] {path.name}: no cached analysis (--no-llm); placeholder slide")
            analyses.append((path, placeholder(path, "")))
            continue
        log(f"[{i}/{len(images)}] {path.name}: asking {model}")
        try:
            analysis = analyse_plot(client, model, path, args.context)
        except Exception as e:  # report and carry on; never abort the whole run
            failures += 1
            log(f"  FAILED: {type(e).__name__}: {e}")
            analysis = placeholder(path, f"Analysis failed: {type(e).__name__}")
            analysis["notes"] = str(e)
        else:
            sidecar["plots"][path.name] = {"hash": h, "model": model, "prompt": PROMPT_VERSION,
                                           "analysis": analysis}
            save_sidecar(sidecar_path, sidecar)   # progress survives a crash mid-run
        analyses.append((path, analysis))

    good = [(p.name, a) for p, a in analyses if not a.get("error")]
    synth_hash = hashlib.sha256(json.dumps(good, sort_keys=True).encode()).hexdigest()[:16]
    cached = sidecar.get("synthesis") or {}
    synth = None
    if args.no_llm:
        synth = cached.get("result")
    elif (cached.get("hash") == synth_hash and cached.get("model") == model
          and cached.get("prompt") == PROMPT_VERSION and not args.no_cache):
        synth = cached["result"]
    elif client is not None and good:
        log(f"synthesising key findings from {len(good)} plots")
        try:
            synth = synthesise(client, model, good, args.context)
            sidecar["synthesis"] = {"hash": synth_hash, "model": model, "prompt": PROMPT_VERSION,
                                    "result": synth}
            save_sidecar(sidecar_path, sidecar)
        except Exception as e:
            log(f"  synthesis FAILED: {type(e).__name__}: {e}")
    synth = synth or {}

    n = build_deck(out, args.title, args.subtitle or synth.get("subtitle", ""),
                   synth.get("key_findings", []), analyses, model)
    log(f"wrote {out}: {n} slides from {len(analyses)} plots ({failures} failed); analyses in {sidecar_path}")


if __name__ == "__main__":
    main()
